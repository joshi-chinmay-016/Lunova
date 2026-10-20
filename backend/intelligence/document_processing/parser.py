"""Generic provider-independent attachment parser for PDF, DOCX, and TXT files.

Extracts text, preserves structural metadata (pages, tables, sections), detects
empty/image-only documents, and reports explicit errors without crashing or leaking
confidential data into logs.
"""

from dataclasses import dataclass, field
from enum import Enum
import io
import os
from pathlib import Path
from typing import Any, Dict, List, Optional, Union
import zipfile

# Try importing pypdf
try:
    import pypdf
    _PYPDF_AVAILABLE = True
except ImportError:
    pypdf = None  # type: ignore
    _PYPDF_AVAILABLE = False

# Try importing docx
try:
    import docx
    _DOCX_AVAILABLE = True
except ImportError:
    docx = None  # type: ignore
    _DOCX_AVAILABLE = False


class DocumentParseStatus(str, Enum):
    """Extraction status classification."""

    SUCCESS = "SUCCESS"
    EMPTY = "EMPTY"
    FAILED = "FAILED"
    UNSUPPORTED = "UNSUPPORTED"
    OVERSIZED = "OVERSIZED"
    CORRUPTED = "CORRUPTED"


DEFAULT_MAX_ATTACHMENT_BYTES = 25 * 1024 * 1024  # 25 MB


@dataclass
class DocumentPage:
    """Individual extracted page or section within a document."""

    page_number: int
    text: str
    metadata: Dict[str, Any] = field(default_factory=dict)


@dataclass
class ParsedDocument:
    """Normalized extraction result from an attachment or file."""

    filename: str
    content_type: str
    text: str
    status: DocumentParseStatus
    pages: List[DocumentPage] = field(default_factory=list)
    page_count: int = 0
    character_count: int = 0
    word_count: int = 0
    error_message: Optional[str] = None
    metadata: Dict[str, Any] = field(default_factory=dict)

    @property
    def is_usable(self) -> bool:
        """Returns True if the document text was successfully extracted and is non-empty."""
        return self.status == DocumentParseStatus.SUCCESS and bool(self.text.strip())


class AttachmentParser:
    """Generic, provider-independent attachment parsing engine.

    Supports:
    - PDF (native text extraction, page tracking, image-only detection)
    - DOCX (paragraphs, tables, basic structure)
    - TXT (multi-encoding detection, line normalization)
    """

    def __init__(self, max_bytes: int = DEFAULT_MAX_ATTACHMENT_BYTES) -> None:
        self.max_bytes = max_bytes

    def parse(
        self,
        content: Union[bytes, io.BytesIO, str, Path],
        filename: str = "document",
        content_type: Optional[str] = None,
    ) -> ParsedDocument:
        """Parse raw document content into a normalized ParsedDocument.

        Args:
            content: Raw file bytes, BytesIO stream, file path, or string.
            filename: Original filename (used for extension detection & reporting).
            content_type: Optional MIME type (e.g. 'application/pdf').

        Returns:
            ParsedDocument containing extracted text, pages, and metadata.
        """
        # 1. Resolve raw bytes
        raw_bytes: bytes
        if isinstance(content, (str, Path)) and os.path.isfile(str(content)):
            filename = os.path.basename(str(content))
            with open(content, "rb") as f:
                raw_bytes = f.read()
        elif isinstance(content, io.BytesIO):
            raw_bytes = content.getvalue()
        elif isinstance(content, bytes):
            raw_bytes = content
        elif isinstance(content, str):
            # Direct text string provided as content
            raw_bytes = content.encode("utf-8")
            if not content_type:
                content_type = "text/plain"
        else:
            return ParsedDocument(
                filename=filename,
                content_type=content_type or "application/octet-stream",
                text="",
                status=DocumentParseStatus.FAILED,
                error_message=f"Unsupported content type input: {type(content).__name__}",
            )

        # 2. Validate file size
        file_size = len(raw_bytes)
        if file_size > self.max_bytes:
            return ParsedDocument(
                filename=filename,
                content_type=content_type or "application/octet-stream",
                text="",
                status=DocumentParseStatus.OVERSIZED,
                error_message=(
                    f"File size ({file_size} bytes) exceeds configured limit ({self.max_bytes} bytes)."
                ),
                metadata={"size_bytes": file_size, "limit_bytes": self.max_bytes},
            )

        # 3. Detect format via magic bytes and content_type / filename
        detected_format, detected_mime = self._detect_format(raw_bytes, filename, content_type)
        resolved_mime = content_type or detected_mime

        if detected_format == "empty":
            return ParsedDocument(
                filename=filename,
                content_type=resolved_mime,
                text="",
                status=DocumentParseStatus.EMPTY,
                error_message="Attachment is completely empty (0 bytes).",
                metadata={"size_bytes": 0},
            )

        # 4. Route to specific format extractor
        try:
            if detected_format == "pdf":
                return self._parse_pdf(raw_bytes, filename, resolved_mime)
            elif detected_format == "docx":
                return self._parse_docx(raw_bytes, filename, resolved_mime)
            elif detected_format == "txt":
                return self._parse_txt(raw_bytes, filename, resolved_mime)
            else:
                return ParsedDocument(
                    filename=filename,
                    content_type=resolved_mime,
                    text="",
                    status=DocumentParseStatus.UNSUPPORTED,
                    error_message=f"Unsupported document format '{detected_format}'. Supported formats: PDF, DOCX, TXT.",
                    metadata={"detected_format": detected_format, "size_bytes": file_size},
                )
        except Exception as exc:
            return ParsedDocument(
                filename=filename,
                content_type=resolved_mime,
                text="",
                status=DocumentParseStatus.FAILED,
                error_message=f"Unexpected extraction failure: {type(exc).__name__}: {str(exc)}",
                metadata={"size_bytes": file_size},
            )

    def _detect_format(
        self,
        raw_bytes: bytes,
        filename: str,
        content_type: Optional[str] = None,
    ) -> tuple[str, str]:
        """Detect file format based on magic bytes, MIME type, and filename extension."""
        if not raw_bytes:
            return "empty", "application/octet-stream"

        lower_fn = filename.lower()
        lower_mime = (content_type or "").lower()

        # Check PDF magic bytes (%PDF-)
        if raw_bytes.startswith(b"%PDF-"):
            return "pdf", "application/pdf"

        # Check DOCX / ZIP magic bytes (PK\x03\x04)
        if raw_bytes.startswith(b"PK\x03\x04"):
            # Could be docx or another zip archive
            if lower_fn.endswith(".docx") or "wordprocessingml" in lower_mime:
                return "docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            # Inspect internal zip structure for word/document.xml
            try:
                with zipfile.ZipFile(io.BytesIO(raw_bytes)) as zf:
                    if "word/document.xml" in zf.namelist():
                        return "docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            except Exception:
                pass

        # Check text MIME or extension
        if (
            lower_fn.endswith((".txt", ".md", ".csv", ".json", ".log", ".rtf"))
            or "text/" in lower_mime
            or lower_mime in ("application/json", "text/plain", "text/markdown")
        ):
            return "txt", "text/plain"

        # Check if bytes are valid UTF-8/ASCII text without null bytes
        if b"\x00" not in raw_bytes[:1024]:
            try:
                raw_bytes[:2048].decode("utf-8")
                return "txt", "text/plain"
            except UnicodeDecodeError:
                pass

        # Fallback to extension check
        if lower_fn.endswith(".pdf"):
            return "pdf", "application/pdf"
        elif lower_fn.endswith(".docx"):
            return "docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

        return "unknown", "application/octet-stream"

    def _parse_pdf(self, raw_bytes: bytes, filename: str, content_type: str) -> ParsedDocument:
        """Extract text and page boundaries from a PDF document."""
        if not _PYPDF_AVAILABLE or pypdf is None:
            return ParsedDocument(
                filename=filename,
                content_type=content_type,
                text="",
                status=DocumentParseStatus.FAILED,
                error_message="pypdf library is not installed in the environment.",
            )

        try:
            stream = io.BytesIO(raw_bytes)
            reader = pypdf.PdfReader(stream)

            if reader.is_encrypted:
                try:
                    # Attempt decrypt with empty password
                    reader.decrypt("")
                except Exception:
                    return ParsedDocument(
                        filename=filename,
                        content_type=content_type,
                        text="",
                        status=DocumentParseStatus.FAILED,
                        error_message="PDF is encrypted and password protected.",
                    )

            pages: List[DocumentPage] = []
            extracted_texts: List[str] = []

            for idx, page in enumerate(reader.pages):
                page_text = page.extract_text() or ""
                cleaned_text = page_text.strip()
                pages.append(
                    DocumentPage(
                        page_number=idx + 1,
                        text=cleaned_text,
                        metadata={"page_index": idx},
                    )
                )
                if cleaned_text:
                    extracted_texts.append(cleaned_text)

            full_text = "\n\n".join(extracted_texts).strip()

            # Detect image-only / empty PDFs
            if not full_text:
                return ParsedDocument(
                    filename=filename,
                    content_type=content_type,
                    text="",
                    status=DocumentParseStatus.EMPTY,
                    pages=pages,
                    page_count=len(reader.pages),
                    error_message=(
                        f"PDF contains {len(reader.pages)} page(s) but no extractable text was found "
                        "(document may be scanned images or empty; OCR is not supported in this phase)."
                    ),
                    metadata={"page_count": len(reader.pages), "is_image_only": True},
                )

            words = len(full_text.split())
            return ParsedDocument(
                filename=filename,
                content_type=content_type,
                text=full_text,
                status=DocumentParseStatus.SUCCESS,
                pages=pages,
                page_count=len(reader.pages),
                character_count=len(full_text),
                word_count=words,
                metadata={
                    "page_count": len(reader.pages),
                    "pages_with_text": len(extracted_texts),
                },
            )
        except Exception as exc:
            return ParsedDocument(
                filename=filename,
                content_type=content_type,
                text="",
                status=DocumentParseStatus.CORRUPTED,
                error_message=f"Failed to parse corrupted PDF: {str(exc)}",
                metadata={"size_bytes": len(raw_bytes)},
            )

    def _parse_docx(self, raw_bytes: bytes, filename: str, content_type: str) -> ParsedDocument:
        """Extract paragraphs and tables from a DOCX document."""
        # Try python-docx first if available
        if _DOCX_AVAILABLE and docx is not None:
            try:
                stream = io.BytesIO(raw_bytes)
                doc = docx.Document(stream)

                elements: List[str] = []

                # Extract paragraphs
                for p in doc.paragraphs:
                    p_text = p.text.strip()
                    if p_text:
                        elements.append(p_text)

                # Extract tables
                for t_idx, table in enumerate(doc.tables):
                    table_rows: List[str] = []
                    for row in table.rows:
                        row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if row_cells:
                            table_rows.append(" | ".join(row_cells))
                    if table_rows:
                        table_text = f"[Table {t_idx + 1}]\n" + "\n".join(table_rows)
                        elements.append(table_text)

                full_text = "\n\n".join(elements).strip()

                if not full_text:
                    return ParsedDocument(
                        filename=filename,
                        content_type=content_type,
                        text="",
                        status=DocumentParseStatus.EMPTY,
                        error_message="DOCX document contains no readable text content.",
                        metadata={"paragraphs_count": len(doc.paragraphs), "tables_count": len(doc.tables)},
                    )

                return ParsedDocument(
                    filename=filename,
                    content_type=content_type,
                    text=full_text,
                    status=DocumentParseStatus.SUCCESS,
                    page_count=1,  # DOCX has no fixed page boundaries
                    character_count=len(full_text),
                    word_count=len(full_text.split()),
                    metadata={"paragraphs": len(doc.paragraphs), "tables": len(doc.tables)},
                )
            except Exception as exc:
                # If python-docx fails or zip is corrupted
                return ParsedDocument(
                    filename=filename,
                    content_type=content_type,
                    text="",
                    status=DocumentParseStatus.CORRUPTED,
                    error_message=f"Failed to parse corrupted DOCX: {str(exc)}",
                    metadata={"size_bytes": len(raw_bytes)},
                )

        # Fallback to direct XML parsing via zipfile (zero external dependencies)
        return self._parse_docx_xml_fallback(raw_bytes, filename, content_type)

    def _parse_docx_xml_fallback(self, raw_bytes: bytes, filename: str, content_type: str) -> ParsedDocument:
        """Parse DOCX word/document.xml directly without external dependencies."""
        import xml.etree.ElementTree as ET

        try:
            with zipfile.ZipFile(io.BytesIO(raw_bytes)) as zf:
                if "word/document.xml" not in zf.namelist():
                    return ParsedDocument(
                        filename=filename,
                        content_type=content_type,
                        text="",
                        status=DocumentParseStatus.CORRUPTED,
                        error_message="Invalid DOCX archive: missing word/document.xml",
                    )
                xml_content = zf.read("word/document.xml")

            tree = ET.fromstring(xml_content)
            # Namespace for WordprocessingML
            ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}

            paragraphs: List[str] = []
            for p in tree.iter(f"{{{ns['w']}}}p"):
                texts = [node.text for node in p.iter(f"{{{ns['w']}}}t") if node.text]
                p_text = "".join(texts).strip()
                if p_text:
                    paragraphs.append(p_text)

            full_text = "\n\n".join(paragraphs).strip()
            if not full_text:
                return ParsedDocument(
                    filename=filename,
                    content_type=content_type,
                    text="",
                    status=DocumentParseStatus.EMPTY,
                    error_message="DOCX document contains no text.",
                )

            return ParsedDocument(
                filename=filename,
                content_type=content_type,
                text=full_text,
                status=DocumentParseStatus.SUCCESS,
                page_count=1,
                character_count=len(full_text),
                word_count=len(full_text.split()),
                metadata={"xml_fallback": True, "paragraphs_count": len(paragraphs)},
            )
        except Exception as exc:
            return ParsedDocument(
                filename=filename,
                content_type=content_type,
                text="",
                status=DocumentParseStatus.CORRUPTED,
                error_message=f"Failed to parse DOCX archive: {str(exc)}",
            )

    def _parse_txt(self, raw_bytes: bytes, filename: str, content_type: str) -> ParsedDocument:
        """Decode and normalize plain text document."""
        encodings = ["utf-8", "utf-8-sig", "latin-1", "cp1252", "iso-8859-1", "utf-16"]
        decoded_text: Optional[str] = None
        used_encoding = "utf-8"

        for enc in encodings:
            try:
                decoded_text = raw_bytes.decode(enc)
                used_encoding = enc
                break
            except (UnicodeDecodeError, LookupError):
                continue

        if decoded_text is None:
            # Fallback with replacement
            decoded_text = raw_bytes.decode("utf-8", errors="replace")
            used_encoding = "utf-8-replace"

        # Normalize line endings
        normalized = decoded_text.replace("\r\n", "\n").replace("\r", "\n").strip()

        if not normalized:
            return ParsedDocument(
                filename=filename,
                content_type=content_type,
                text="",
                status=DocumentParseStatus.EMPTY,
                error_message="Text document is empty or contains only whitespace.",
                metadata={"encoding": used_encoding, "size_bytes": len(raw_bytes)},
            )

        return ParsedDocument(
            filename=filename,
            content_type=content_type,
            text=normalized,
            status=DocumentParseStatus.SUCCESS,
            page_count=1,
            character_count=len(normalized),
            word_count=len(normalized.split()),
            metadata={"encoding": used_encoding, "size_bytes": len(raw_bytes)},
        )
